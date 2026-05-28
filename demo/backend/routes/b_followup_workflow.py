"""
B端随访任务与AI企微工作流 API
"""
import base64
from datetime import datetime, timedelta
import hashlib
import re
import struct
import xml.etree.ElementTree as ET
import uuid

from cryptography.hazmat.backends import default_backend
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from flask import Blueprint, request, g, current_app

from models import (
    db,
    User,
    BPatient,
    BHealthRecord,
    BReport,
    BFollowUpTask,
    BFollowUpMessage,
    BFollowUpEvent,
    BFollowUpKnowledgeItem,
    BFollowUpPlanTemplate,
    BFollowUpPlanNode,
    BFollowUpAIRule,
    BFollowUpPatientPlan,
    BFollowUpCheckin,
)
from services.followup_ai_service import followup_ai_service
from services.followup_checkin_service import followup_checkin_service
from services.followup_scheduler_service import followup_scheduler_service
from services.wecom_service import wecom_service
from utils.decorators import login_required
from utils.response import Response

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - import error is handled at runtime
    load_workbook = None


b_followup_bp = Blueprint('b_followup', __name__, url_prefix='/api/b/followup')
wecom_callback_bp = Blueprint('wecom_callback', __name__, url_prefix='/api/wecom')


@b_followup_bp.route('/scheduler/run-once', methods=['POST'])
@login_required
def run_scheduler_once():
    """手动触发一次随访任务调度。生产可由定时器调用同一服务。"""
    data = request.json or {}
    result = followup_scheduler_service.run_once(limit=int(data.get('limit') or 50))
    return Response.success(result, '随访调度已执行')


def _parse_datetime(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try:
            parsed = datetime.strptime(text, fmt)
            if fmt == '%Y-%m-%d':
                parsed = parsed.replace(hour=9, minute=0, second=0)
            return parsed
        except ValueError:
            continue
    raise ValueError('时间格式错误，请使用 YYYY-MM-DD 或 YYYY-MM-DD HH:mm:ss')


def _task_code():
    return f"FU{datetime.now().strftime('%Y%m%d%H%M%S')}{uuid.uuid4().hex[:6].upper()}"


def _public_checkin_url(task):
    """Return the patient-facing check-in URL for a task."""
    base = (
        current_app.config.get('FOLLOWUP_PUBLIC_BASE_URL')
        or current_app.config.get('PUBLIC_BASE_URL')
        or request.host_url.rstrip('/')
    )
    return f"{base}/followup-checkin/{task.task_code}"


def _recipient_for_channel(patient, channel):
    if not patient:
        return None
    if channel == 'wecom':
        return patient.wecom_external_userid or patient.wecom_userid or patient.wechat_id
    if channel == 'phone':
        return patient.phone
    if channel == 'miniapp':
        return patient.wechat_id
    return patient.phone or patient.wecom_external_userid or patient.wecom_userid or patient.wechat_id


def _append_checkin_link(content, checkin_url):
    if not checkin_url or checkin_url in content:
        return content
    suffix = f"\n\n请通过以下链接完成本次健康打卡：\n{checkin_url}"
    return f"{content.rstrip()}{suffix}" if content else suffix.strip()


def _dispatch_followup_message(task, content):
    """
    Dispatch a follow-up message by channel.

    Non-WeCom channels intentionally produce a dry-run record with the public
    check-in URL, so the task can be completed without enterprise WeChat.
    """
    checkin_url = _public_checkin_url(task)
    content_with_link = _append_checkin_link(content, checkin_url)
    channel = task.channel or 'manual'

    if channel == 'wecom' and task.channel_recipient:
        result = wecom_service.send_text(task.channel_recipient, content_with_link)
    else:
        result = {
            'ok': True,
            'dry_run': True,
            'status': 'dry_run',
            'channel': channel,
            'recipient': task.channel_recipient,
            'payload': {
                'channel': channel,
                'recipient': task.channel_recipient,
                'content': content_with_link,
                'checkin_url': checkin_url,
            },
            'message': '未启用真实外部触达或缺少接收人，已生成公开打卡链接并保存发送记录',
        }

    result['content'] = content_with_link
    result['checkin_url'] = checkin_url
    return result


def _plan_code():
    return f"FP{datetime.now().strftime('%Y%m%d%H%M%S')}{uuid.uuid4().hex[:6].upper()}"


def _template_code():
    return f"FT{datetime.now().strftime('%Y%m%d%H%M%S')}{uuid.uuid4().hex[:6].upper()}"


def _excel_import_template_code():
    return 'TPL_THYROID_LUNG_90D_MENTAL_EXCEL'


def _excel_time_to_hhmm(value):
    text = str(value or '').strip().replace(' ', '')
    mapping = {
        '上午七点': '07:00',
        '早上七点': '07:00',
        '中午十一点': '11:00',
        '上午十一点': '11:00',
        '中午十二点半': '12:30',
        '晚上五点半': '17:30',
        '下午五点半': '17:30',
        '晚上七点': '19:00',
        '晚上八点': '20:00',
    }
    if text in mapping:
        return mapping[text]
    match = re.search(r'(\d{1,2})[:：](\d{1,2})', text)
    if match:
        return f"{int(match.group(1)):02d}:{int(match.group(2)):02d}"
    return '09:00'


def _excel_node_kind(title, time_text=''):
    text = f"{title or ''} {time_text or ''}"
    if '知识' in text:
        return {
            'task_type': 'knowledge_push',
            'patient_action': 'none',
            'ai_action': 'reply',
            'label': '知识卡'
        }
    if '运动' in text:
        return {
            'task_type': 'exercise_reminder',
            'patient_action': 'reply_text',
            'ai_action': 'reply',
            'label': '运动提醒'
        }
    if '心理' in text:
        return {
            'task_type': 'psych_reminder',
            'patient_action': 'reply_text',
            'ai_action': 'reply',
            'label': '心理关怀'
        }
    if '早餐' in text or '早起' in text:
        label = '早餐打卡'
    elif '午餐' in text:
        label = '午餐打卡'
    elif '晚餐' in text:
        label = '晚餐打卡'
    else:
        label = str(title or '').strip() or '饮食打卡'
    return {
        'task_type': 'diet_checkin',
        'patient_action': 'upload_image',
        'ai_action': 'diet_review',
        'label': label
    }


def _parse_followup_excel(file_storage):
    if load_workbook is None:
        raise RuntimeError('服务器缺少 openpyxl，无法解析 Excel 文件')
    workbook = load_workbook(file_storage, read_only=True, data_only=True)
    nodes = []
    parsed_days = 0
    for sheet_name in workbook.sheetnames:
        match = re.match(r'^day\s*(\d+)$', str(sheet_name).strip(), re.I)
        if not match:
            continue
        day = int(match.group(1))
        ws = workbook[sheet_name]
        intro = ''
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if first_row and first_row[0]:
            intro = str(first_row[0]).strip()
        sort_index = 0
        for row in ws.iter_rows(min_row=3, values_only=True):
            time_text = str(row[0] or '').strip() if len(row) > 0 else ''
            title = str(row[1] or '').strip() if len(row) > 1 else ''
            content = str(row[2] or '').strip() if len(row) > 2 else ''
            if not time_text or not title or not content:
                continue
            kind = _excel_node_kind(title, time_text)
            message = content
            if sort_index == 0 and intro:
                message = f"{intro}\n\n{content}"
            send_time = _excel_time_to_hhmm(time_text)
            nodes.append({
                'node_code': f"D{day:03d}_{sort_index + 1:02d}_{kind['task_type'].upper()}",
                'name': f"Day {day} {kind['label']}",
                'day_offset': day,
                'send_time': send_time,
                'task_type': kind['task_type'],
                'patient_action': kind['patient_action'],
                'ai_action': kind['ai_action'],
                'message_template': message,
                'knowledge_item_ids': [],
                'checkin_schema': {
                    'source': 'excel_import',
                    'excel_sheet': sheet_name,
                    'excel_time': time_text,
                    'excel_title': title,
                },
                'escalation_rule': {
                    'manual_keywords': ['焦虑', '睡不着', '担心', '不适'],
                    'doctor_keywords': ['疼痛', '明显加重', '胸闷', '呼吸困难'],
                    'no_reply_threshold': 3,
                    'abnormal_actions': {
                        'manual_keywords': 'manual_handoff',
                        'doctor_keywords': 'doctor_handoff',
                        'no_reply': 'manual_handoff'
                    }
                },
                'completion_rule': {
                    'type': 'checkin_submitted' if kind['patient_action'] == 'upload_image' else 'message_sent',
                    'overdue_action': 'manual_handoff',
                    'required': True,
                },
                'is_required': True,
                'is_active': True,
                'sort_order': (day * 100) + sort_index,
            })
            sort_index += 1
        if sort_index:
            parsed_days += 1
    if not nodes:
        raise ValueError('Excel 未解析到任务节点，请确认工作表为 day1-day90，且包含“时间/概述/温馨提醒”列')
    nodes.sort(key=lambda item: (item['day_offset'], item['sort_order']))
    return {'days': parsed_days, 'nodes': nodes}


def _event(task, event_type, summary='', from_status=None, to_status=None, actor_type='system', payload=None):
    item = BFollowUpEvent(
        task_id=task.id,
        event_type=event_type,
        from_status=from_status,
        to_status=to_status,
        actor_type=actor_type,
        actor_id=str(getattr(g, 'user_id', '') or ''),
        summary=summary,
        payload=payload or {}
    )
    db.session.add(item)
    return item


def _fallback_manager_id(patient=None):
    if patient and patient.manager_id:
        return patient.manager_id
    user = User.query.filter_by(is_active=True).order_by(User.id.asc()).first()
    return user.id if user else None


def _latest_record(patient_id):
    return BHealthRecord.query.filter_by(patient_id=patient_id).order_by(BHealthRecord.created_at.desc()).first()


def _latest_report(patient_id, record_id=None):
    query = BReport.query.filter_by(patient_id=patient_id)
    if record_id:
        query = query.filter_by(record_id=record_id)
    return query.order_by(BReport.created_at.desc()).first()


def _normalize_risk_for_match(value):
    text = str(value or '').strip().lower()
    if text in ('high', '高危', '高风险', '高'):
        return 'high'
    if text in ('mid', 'medium', '中危', '中风险', '中'):
        return 'mid'
    if text in ('low', '低危', '低风险', '低'):
        return 'low'
    return text or None


def _risk_aliases(value):
    risk = _normalize_risk_for_match(value)
    if risk == 'high':
        return ['high', '高危', '高风险']
    if risk == 'mid':
        return ['mid', 'medium', '中危', '中风险']
    if risk == 'low':
        return ['low', '低危', '低风险']
    return [value] if value else []


def _first_followup_days(risk_level):
    risk = _normalize_risk_for_match(risk_level)
    if risk == 'high':
        return 7
    if risk == 'mid':
        return 30
    if risk == 'low':
        return 90
    return 30


def _open_report_followup_task(report_id):
    return (
        BFollowUpTask.query
        .filter(
            BFollowUpTask.report_id == report_id,
            BFollowUpTask.source == 'report',
            BFollowUpTask.status != 'cancelled'
        )
        .order_by(BFollowUpTask.created_at.desc())
        .first()
    )


def _match_knowledge(patient=None, nodule_type=None, risk_level=None, task_type=None, categories=None, limit=20):
    query = BFollowUpKnowledgeItem.query.filter(BFollowUpKnowledgeItem.is_active.is_(True))
    if categories:
        query = query.filter(BFollowUpKnowledgeItem.category.in_(categories))
    if nodule_type:
        query = query.filter(db.or_(
            BFollowUpKnowledgeItem.nodule_type.is_(None),
            BFollowUpKnowledgeItem.nodule_type == '',
            BFollowUpKnowledgeItem.nodule_type == nodule_type
        ))
    if risk_level:
        aliases = _risk_aliases(risk_level)
        query = query.filter(db.or_(
            BFollowUpKnowledgeItem.risk_level.is_(None),
            BFollowUpKnowledgeItem.risk_level == '',
            BFollowUpKnowledgeItem.risk_level.in_(aliases)
        ))
    if task_type:
        query = query.filter(db.or_(
            BFollowUpKnowledgeItem.task_type.is_(None),
            BFollowUpKnowledgeItem.task_type == '',
            BFollowUpKnowledgeItem.task_type == task_type
        ))
    if patient and patient.age:
        query = query.filter(
            (BFollowUpKnowledgeItem.age_min.is_(None) | (BFollowUpKnowledgeItem.age_min <= patient.age)) &
            (BFollowUpKnowledgeItem.age_max.is_(None) | (BFollowUpKnowledgeItem.age_max >= patient.age))
        )
    if patient and patient.gender:
        query = query.filter(db.or_(
            BFollowUpKnowledgeItem.gender.is_(None),
            BFollowUpKnowledgeItem.gender == '',
            BFollowUpKnowledgeItem.gender == patient.gender
        ))
    return query.order_by(BFollowUpKnowledgeItem.priority.desc(), BFollowUpKnowledgeItem.id.desc()).limit(limit).all()


def _seed_default_followup_config(user_id=None):
    """创建一组最小可用的任务提醒、知识库和餐饮识别配置。幂等。"""
    if BFollowUpPlanTemplate.query.filter_by(template_code='TPL_HEALTH_TASK_DAILY').first():
        return False

    defaults = [
        BFollowUpKnowledgeItem(
            title='肺结节日常管理知识',
            category='lung_nodule',
            content='肺结节用户日常管理以按期复查、避免吸烟和二手烟、规律作息、记录咳嗽胸闷等变化为重点。内容用于健康管理提醒，不替代医生诊断。',
            nodule_type='lung',
            task_type='knowledge_push',
            trigger_keywords='肺结节,复查,戒烟,咳嗽',
            action_type='push',
            priority=10,
            source='system_seed'
        ),
        BFollowUpKnowledgeItem(
            title='饮食打卡与餐饮图片分析规则',
            category='diet',
            content='餐饮分析关注主食、蔬菜、优质蛋白、油脂、糖分、饮酒和加工食品。回复应给出可执行的下一餐建议，不将饮食直接归因为结节变化。',
            task_type='diet_checkin',
            action_type='diet_review',
            priority=9,
            source='system_seed'
        ),
        BFollowUpKnowledgeItem(
            title='运动提醒规则',
            category='exercise',
            content='运动提醒以低到中等强度、持续性、可坚持为原则。建议散步、八段锦、轻力量训练等，避免突然大强度运动。',
            task_type='exercise_reminder',
            action_type='push',
            priority=8,
            source='system_seed'
        ),
        BFollowUpKnowledgeItem(
            title='心理与睡眠提醒',
            category='psych',
            content='心理提醒关注焦虑、睡眠、压力和对检查结果的担忧。建议使用安抚性语言，引导规律作息、呼吸放松和必要时联系健康管理师。',
            task_type='psych_reminder',
            trigger_keywords='焦虑,睡不着,担心,压力',
            action_type='push',
            priority=7,
            source='system_seed'
        ),
        BFollowUpKnowledgeItem(
            title='乳腺结节健康知识',
            category='breast_nodule',
            content='乳腺结节用户关注按期复查、乳房自我观察、规律作息、减少焦虑和保持适量运动。出现明显疼痛、红肿、溢液等情况应及时咨询医生。',
            nodule_type='breast',
            task_type='knowledge_push',
            trigger_keywords='乳腺结节,乳房,复查',
            action_type='push',
            priority=7,
            source='system_seed'
        ),
        BFollowUpKnowledgeItem(
            title='甲状腺结节健康知识',
            category='thyroid_nodule',
            content='甲状腺结节用户关注按期超声复查、情绪压力管理、规律睡眠和遵医嘱检查甲状腺功能。',
            nodule_type='thyroid',
            task_type='knowledge_push',
            trigger_keywords='甲状腺结节,超声,情绪',
            action_type='push',
            priority=7,
            source='system_seed'
        ),
    ]
    db.session.add_all(defaults)
    db.session.flush()

    template = BFollowUpPlanTemplate(
        template_code='TPL_HEALTH_TASK_DAILY',
        name='健康管理每日任务模板',
        description='面向结节健康管理用户，覆盖结节知识、饮食、运动、心理提醒和餐饮图片识别。',
        nodule_type=None,
        risk_level=None,
        cycle_days=30,
        default_channel='wecom',
        default_reminder_strategy='每日固定时间提醒；未打卡继续提醒；餐饮图片提交后自动分析并记录。',
        audience_rule={'module': 'task_reminder_and_diet_image_analysis'},
        status='active',
        version=1,
        created_by=user_id
    )
    db.session.add(template)
    db.session.flush()

    nodes = [
        BFollowUpPlanNode(
            template_id=template.id,
            node_code='DAILY_NODULE_KNOWLEDGE',
            name='结节知识推送',
            day_offset=1,
            send_time='08:30',
            task_type='knowledge_push',
            patient_action='none',
            ai_action='reply',
            message_template='早上好，今天为您推送一条结节健康管理知识。请按计划保持规律作息、适量运动并记录身体变化。',
            knowledge_item_ids=[defaults[0].id, defaults[4].id, defaults[5].id],
            escalation_rule={},
            completion_rule={'type': 'message_sent'}
        ),
        BFollowUpPlanNode(
            template_id=template.id,
            node_code='DAILY_CHECKIN',
            name='每日健康打卡提醒',
            day_offset=1,
            send_time='09:00',
            task_type='daily_checkin',
            patient_action='reply_text',
            ai_action='reply',
            message_template='请完成今日健康打卡：睡眠、饮食、运动和情绪状态。如有不适，也可以在备注中说明。',
            knowledge_item_ids=[defaults[3].id],
            checkin_schema={'fields': ['睡眠', '饮食', '运动', '情绪', '备注']},
            escalation_rule={'manual_keywords': ['焦虑', '睡不着', '不舒服']},
            completion_rule={'type': 'checkin_submitted'}
        ),
        BFollowUpPlanNode(
            template_id=template.id,
            node_code='DIET_IMAGE_CHECKIN',
            name='餐饮图片识别打卡',
            day_offset=1,
            send_time='12:30',
            task_type='diet_checkin',
            patient_action='upload_image',
            ai_action='diet_review',
            message_template='请上传今天一餐照片，系统会从主食、蔬菜、蛋白质、油脂和糖分角度给出饮食建议。',
            knowledge_item_ids=[defaults[1].id],
            checkin_schema={'image_required': True, 'fields': ['餐别', '饱腹感', '饮酒/甜饮', '备注']},
            escalation_rule={'manual_if': '连续多次高油高糖或饮酒'},
            completion_rule={'type': 'checkin_submitted'}
        ),
        BFollowUpPlanNode(
            template_id=template.id,
            node_code='EXERCISE_REMINDER',
            name='运动提醒',
            day_offset=1,
            send_time='18:30',
            task_type='exercise_reminder',
            patient_action='none',
            ai_action='reply',
            message_template='今天建议安排20-30分钟低到中等强度活动，如散步、拉伸或八段锦。量力而行，贵在坚持。',
            knowledge_item_ids=[defaults[2].id],
            escalation_rule={},
            completion_rule={'type': 'message_sent'}
        ),
        BFollowUpPlanNode(
            template_id=template.id,
            node_code='PSYCH_REMINDER',
            name='心理睡眠提醒',
            day_offset=1,
            send_time='21:00',
            task_type='psych_reminder',
            patient_action='reply_text',
            ai_action='reply',
            message_template='睡前可以做3分钟呼吸放松，记录今天的压力和睡眠准备情况。若持续焦虑或失眠，可回复说明。',
            knowledge_item_ids=[defaults[3].id],
            escalation_rule={'manual_keywords': ['焦虑', '睡不着', '害怕', '压力大']},
            completion_rule={'type': 'patient_reply'}
        ),
    ]
    db.session.add_all(nodes)

    rules = [
        BFollowUpAIRule(
            name='连续未打卡提醒',
            rule_type='no_reply',
            task_type='daily_checkin',
            trigger_keywords='未打卡,未回复',
            action='notify',
            response_template='您今天的健康打卡还未完成，请在方便时补充睡眠、饮食、运动和情绪情况。',
            priority=10
        ),
        BFollowUpAIRule(
            name='餐饮图片识别点评',
            rule_type='image_recognition',
            task_type='diet_checkin',
            trigger_keywords='高油,高糖,饮酒,蔬菜不足',
            action='ai_reply',
            response_template='这餐我会从主食、蔬菜、优质蛋白、油脂和糖分角度点评，并给出下一餐建议。',
            priority=9
        ),
        BFollowUpAIRule(
            name='心理压力转健康管理师',
            rule_type='escalation',
            task_type='psych_reminder',
            trigger_keywords='焦虑,睡不着,害怕,压力大',
            action='manual_handoff',
            response_template='已收到您的反馈。建议先做简单呼吸放松并保持规律作息，健康管理师会继续跟进您的情况。',
            priority=7
        ),
    ]
    db.session.add_all(rules)
    return True


def _serialize_task_query(query, include_patient=True):
    items = []
    for task in query:
        data = task.to_dict()
        if not include_patient:
            data.pop('patient', None)
        items.append(data)
    return items


@b_followup_bp.route('/config/seed-defaults', methods=['POST'])
@login_required
def seed_default_config():
    """默认模板初始化已关闭，模板需通过Excel导入或手动创建。"""
    return Response.success({'created': False}, '默认模板初始化已关闭，请导入Excel模板或手动创建')


@b_followup_bp.route('/knowledge', methods=['GET'])
@login_required
def list_followup_knowledge():
    """随访知识库列表"""
    page = request.args.get('page', 1, type=int)
    per_page = min(request.args.get('per_page', 20, type=int), 100)
    category = request.args.get('category', '').strip()
    nodule_type = request.args.get('nodule_type', '').strip()
    risk_level = request.args.get('risk_level', '').strip()
    task_type = request.args.get('task_type', '').strip()
    search = request.args.get('search', '').strip()

    query = BFollowUpKnowledgeItem.query
    if category:
        query = query.filter(BFollowUpKnowledgeItem.category == category)
    if nodule_type:
        query = query.filter(BFollowUpKnowledgeItem.nodule_type == nodule_type)
    if risk_level:
        query = query.filter(BFollowUpKnowledgeItem.risk_level.in_(_risk_aliases(risk_level)))
    if task_type:
        query = query.filter(BFollowUpKnowledgeItem.task_type == task_type)
    if search:
        query = query.filter(db.or_(
            BFollowUpKnowledgeItem.title.ilike(f'%{search}%'),
            BFollowUpKnowledgeItem.content.ilike(f'%{search}%')
        ))

    pagination = query.order_by(BFollowUpKnowledgeItem.priority.desc(), BFollowUpKnowledgeItem.id.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    return Response.paginate([item.to_dict() for item in pagination.items], pagination.total, page, per_page)


@b_followup_bp.route('/knowledge', methods=['POST'])
@login_required
def create_followup_knowledge():
    data = request.json or {}
    if not data.get('title') or not data.get('content') or not data.get('category'):
        return Response.error('title、category、content 必填', 400)
    item = BFollowUpKnowledgeItem(
        title=data.get('title'),
        category=data.get('category'),
        content=data.get('content'),
        nodule_type=data.get('nodule_type'),
        risk_level=data.get('risk_level'),
        task_type=data.get('task_type'),
        age_min=data.get('age_min'),
        age_max=data.get('age_max'),
        gender=data.get('gender'),
        trigger_keywords=data.get('trigger_keywords'),
        action_type=data.get('action_type'),
        priority=data.get('priority', 5),
        is_active=data.get('is_active', True),
        source=data.get('source') or 'manual',
        metadata_json=data.get('metadata_json') or {}
    )
    db.session.add(item)
    db.session.commit()
    return Response.success(item.to_dict(), '随访知识条目创建成功', 201)


@b_followup_bp.route('/knowledge/<int:item_id>', methods=['PUT'])
@login_required
def update_followup_knowledge(item_id):
    item = BFollowUpKnowledgeItem.query.get_or_404(item_id)
    data = request.json or {}
    for field in [
        'title', 'category', 'content', 'nodule_type', 'risk_level', 'task_type',
        'age_min', 'age_max', 'gender', 'trigger_keywords', 'action_type',
        'priority', 'is_active', 'source', 'metadata_json'
    ]:
        if field in data:
            setattr(item, field, data[field])
    db.session.commit()
    return Response.success(item.to_dict(), '随访知识条目已更新')


@b_followup_bp.route('/knowledge/<int:item_id>', methods=['DELETE'])
@login_required
def delete_followup_knowledge(item_id):
    item = BFollowUpKnowledgeItem.query.get_or_404(item_id)
    db.session.delete(item)
    db.session.commit()
    return Response.success({'id': item_id}, '随访知识条目已删除')


@b_followup_bp.route('/templates', methods=['GET'])
@login_required
def list_plan_templates():
    """随访计划模板列表"""
    status = request.args.get('status', '').strip()
    nodule_type = request.args.get('nodule_type', '').strip()
    risk_level = request.args.get('risk_level', '').strip()
    include_nodes = request.args.get('include_nodes') == '1'
    query = BFollowUpPlanTemplate.query
    if status:
        query = query.filter(BFollowUpPlanTemplate.status == status)
    if nodule_type:
        query = query.filter(BFollowUpPlanTemplate.nodule_type == nodule_type)
    if risk_level:
        query = query.filter(BFollowUpPlanTemplate.risk_level.in_(_risk_aliases(risk_level)))
    items = query.order_by(BFollowUpPlanTemplate.updated_at.desc()).all()
    return Response.success([item.to_dict(include_nodes=include_nodes) for item in items])


@b_followup_bp.route('/templates', methods=['POST'])
@login_required
def create_plan_template():
    data = request.json or {}
    if not data.get('name'):
        return Response.error('name 必填', 400)
    template = BFollowUpPlanTemplate(
        template_code=data.get('template_code') or _template_code(),
        name=data.get('name'),
        description=data.get('description'),
        nodule_type=data.get('nodule_type'),
        risk_level=data.get('risk_level'),
        cycle_days=data.get('cycle_days') or 90,
        default_channel=data.get('default_channel') or 'wecom',
        default_reminder_strategy=data.get('default_reminder_strategy') or '每日固定时间提醒；未打卡继续提醒；餐饮图片自动分析',
        audience_rule=data.get('audience_rule') or {},
        status=data.get('status') or 'draft',
        version=data.get('version') or 1,
        created_by=g.user_id
    )
    db.session.add(template)
    db.session.flush()
    for idx, node_data in enumerate(data.get('nodes') or []):
        node = BFollowUpPlanNode(
            template_id=template.id,
            node_code=node_data.get('node_code') or f'NODE_{idx + 1}',
            name=node_data.get('name') or f'节点{idx + 1}',
            day_offset=node_data.get('day_offset') or 1,
            send_time=node_data.get('send_time') or '09:00',
            task_type=node_data.get('task_type') or 'knowledge_push',
            patient_action=node_data.get('patient_action') or 'reply_text',
            ai_action=node_data.get('ai_action') or 'reply',
            message_template=node_data.get('message_template') or '请按健康管理任务完成本次打卡或查看提醒。',
            knowledge_item_ids=node_data.get('knowledge_item_ids') or [],
            checkin_schema=node_data.get('checkin_schema') or {},
            escalation_rule=node_data.get('escalation_rule') or {},
            completion_rule=node_data.get('completion_rule') or {},
            is_required=node_data.get('is_required', True),
            is_active=node_data.get('is_active', True),
            sort_order=node_data.get('sort_order', idx)
        )
        db.session.add(node)
    db.session.commit()
    return Response.success(template.to_dict(include_nodes=True), '随访计划模板创建成功', 201)


@b_followup_bp.route('/templates/import-excel', methods=['POST'])
@login_required
def import_plan_template_from_excel():
    """从 90 天 Excel 管理方案导入随访任务模板。"""
    upload = request.files.get('file')
    if not upload:
        return Response.error('请上传 Excel 文件', 400)
    filename = upload.filename or ''
    if not filename.lower().endswith(('.xlsx', '.xlsm')):
        return Response.error('仅支持 .xlsx/.xlsm 文件', 400)

    try:
        parsed = _parse_followup_excel(upload)
    except Exception as e:
        return Response.error(f'Excel解析失败：{e}', 400)

    data = request.form or {}
    base_code = data.get('template_code') or _excel_import_template_code()
    name = data.get('name') or '甲状腺结节合并肺结节90天健康管理模板'
    replace_existing = str(data.get('replace_existing', 'true')).lower() in ('1', 'true', 'yes', 'y')

    template = BFollowUpPlanTemplate.query.filter_by(template_code=base_code).first()
    if template and BFollowUpPatientPlan.query.filter_by(template_id=template.id).first():
        template = None
        base_code = f"{base_code}_{datetime.now().strftime('%Y%m%d%H%M%S')}"

    if template and replace_existing:
        BFollowUpPlanNode.query.filter_by(template_id=template.id).delete(synchronize_session=False)
        template.name = name
        template.description = data.get('description') or '由90天Excel健康管理方案导入，覆盖早餐、午餐、知识卡、晚餐、运动和心理提醒。'
        template.nodule_type = data.get('nodule_type') or 'lung_thyroid'
        template.risk_level = data.get('risk_level') or None
        template.cycle_days = int(data.get('cycle_days') or 90)
        template.default_channel = data.get('default_channel') or 'wecom'
        template.default_reminder_strategy = data.get('default_reminder_strategy') or '每日固定时间企业微信推送；餐饮打卡可上传图片；异常关键词转人工。'
        template.audience_rule = {
            'source': 'excel_import',
            'file_name': filename,
            'nodule_combo': 'lung_thyroid',
            'days': parsed['days'],
            'node_count': len(parsed['nodes']),
        }
        template.status = data.get('status') or 'active'
        template.version = (template.version or 1) + 1
        template.updated_at = datetime.utcnow()
    else:
        template = BFollowUpPlanTemplate(
            template_code=base_code if not template else _template_code(),
            name=name,
            description=data.get('description') or '由90天Excel健康管理方案导入，覆盖早餐、午餐、知识卡、晚餐、运动和心理提醒。',
            nodule_type=data.get('nodule_type') or 'lung_thyroid',
            risk_level=data.get('risk_level') or None,
            cycle_days=int(data.get('cycle_days') or 90),
            default_channel=data.get('default_channel') or 'wecom',
            default_reminder_strategy=data.get('default_reminder_strategy') or '每日固定时间企业微信推送；餐饮打卡可上传图片；异常关键词转人工。',
            audience_rule={
                'source': 'excel_import',
                'file_name': filename,
                'nodule_combo': 'lung_thyroid',
                'days': parsed['days'],
                'node_count': len(parsed['nodes']),
            },
            status=data.get('status') or 'active',
            version=1,
            created_by=g.user_id,
        )
        db.session.add(template)
        db.session.flush()

    for node_data in parsed['nodes']:
        db.session.add(BFollowUpPlanNode(template_id=template.id, **node_data))
    db.session.commit()
    return Response.success({
        'template': template.to_dict(include_nodes=True),
        'days': parsed['days'],
        'nodes': len(parsed['nodes'])
    }, 'Excel模板已导入')


@b_followup_bp.route('/templates/<int:template_id>', methods=['GET'])
@login_required
def get_plan_template(template_id):
    template = BFollowUpPlanTemplate.query.get_or_404(template_id)
    return Response.success(template.to_dict(include_nodes=True))


@b_followup_bp.route('/templates/<int:template_id>', methods=['PUT'])
@login_required
def update_plan_template(template_id):
    template = BFollowUpPlanTemplate.query.get_or_404(template_id)
    data = request.json or {}
    for field in [
        'name', 'description', 'nodule_type', 'risk_level', 'cycle_days',
        'default_channel', 'default_reminder_strategy', 'audience_rule',
        'status', 'version'
    ]:
        if field in data:
            setattr(template, field, data[field])
    db.session.commit()
    return Response.success(template.to_dict(include_nodes=True), '随访计划模板已更新')


@b_followup_bp.route('/templates/<int:template_id>', methods=['DELETE'])
@login_required
def delete_plan_template(template_id):
    template = BFollowUpPlanTemplate.query.get_or_404(template_id)
    if BFollowUpPatientPlan.query.filter_by(template_id=template_id).first():
        return Response.error('该模板已被患者计划引用，不能删除，可将状态改为 archived', 409)
    db.session.delete(template)
    db.session.commit()
    return Response.success({'id': template_id}, '随访计划模板已删除')


@b_followup_bp.route('/templates/<int:template_id>/nodes', methods=['POST'])
@login_required
def create_plan_node(template_id):
    template = BFollowUpPlanTemplate.query.get_or_404(template_id)
    data = request.json or {}
    node = BFollowUpPlanNode(
        template_id=template.id,
        node_code=data.get('node_code') or f'NODE_{uuid.uuid4().hex[:6].upper()}',
        name=data.get('name') or '任务节点',
        day_offset=data.get('day_offset') or 1,
        send_time=data.get('send_time') or '09:00',
        task_type=data.get('task_type') or 'knowledge_push',
        patient_action=data.get('patient_action') or 'reply_text',
        ai_action=data.get('ai_action') or 'reply',
        message_template=data.get('message_template') or '请按健康管理任务完成本次打卡或查看提醒。',
        knowledge_item_ids=data.get('knowledge_item_ids') or [],
        checkin_schema=data.get('checkin_schema') or {},
        escalation_rule=data.get('escalation_rule') or {},
        completion_rule=data.get('completion_rule') or {},
        is_required=data.get('is_required', True),
        is_active=data.get('is_active', True),
        sort_order=data.get('sort_order') or 0
    )
    db.session.add(node)
    db.session.commit()
    return Response.success(node.to_dict(), '随访计划节点创建成功', 201)


@b_followup_bp.route('/nodes/<int:node_id>', methods=['PUT'])
@login_required
def update_plan_node(node_id):
    node = BFollowUpPlanNode.query.get_or_404(node_id)
    data = request.json or {}
    for field in [
        'node_code', 'name', 'day_offset', 'send_time', 'task_type', 'patient_action',
        'ai_action', 'message_template', 'knowledge_item_ids', 'checkin_schema',
        'escalation_rule', 'completion_rule', 'is_required', 'is_active', 'sort_order'
    ]:
        if field in data:
            setattr(node, field, data[field])
    db.session.commit()
    return Response.success(node.to_dict(), '随访计划节点已更新')


@b_followup_bp.route('/nodes/<int:node_id>', methods=['DELETE'])
@login_required
def delete_plan_node(node_id):
    node = BFollowUpPlanNode.query.get_or_404(node_id)
    db.session.delete(node)
    db.session.commit()
    return Response.success({'id': node_id}, '随访计划节点已删除')


@b_followup_bp.route('/ai-rules', methods=['GET'])
@login_required
def list_ai_rules():
    rule_type = request.args.get('rule_type', '').strip()
    task_type = request.args.get('task_type', '').strip()
    query = BFollowUpAIRule.query
    if rule_type:
        query = query.filter(BFollowUpAIRule.rule_type == rule_type)
    if task_type:
        query = query.filter(BFollowUpAIRule.task_type == task_type)
    items = query.order_by(BFollowUpAIRule.priority.desc(), BFollowUpAIRule.id.desc()).all()
    return Response.success([item.to_dict() for item in items])


@b_followup_bp.route('/ai-rules', methods=['POST'])
@login_required
def create_ai_rule():
    data = request.json or {}
    if not data.get('name') or not data.get('rule_type') or not data.get('action'):
        return Response.error('name、rule_type、action 必填', 400)
    rule = BFollowUpAIRule(
        name=data.get('name'),
        rule_type=data.get('rule_type'),
        nodule_type=data.get('nodule_type'),
        risk_level=data.get('risk_level'),
        task_type=data.get('task_type'),
        trigger_keywords=data.get('trigger_keywords'),
        condition_json=data.get('condition_json') or {},
        action=data.get('action'),
        response_template=data.get('response_template'),
        priority=data.get('priority', 5),
        is_active=data.get('is_active', True)
    )
    db.session.add(rule)
    db.session.commit()
    return Response.success(rule.to_dict(), 'AI规则创建成功', 201)


@b_followup_bp.route('/ai-rules/<int:rule_id>', methods=['PUT'])
@login_required
def update_ai_rule(rule_id):
    rule = BFollowUpAIRule.query.get_or_404(rule_id)
    data = request.json or {}
    for field in [
        'name', 'rule_type', 'nodule_type', 'risk_level', 'task_type',
        'trigger_keywords', 'condition_json', 'action', 'response_template',
        'priority', 'is_active'
    ]:
        if field in data:
            setattr(rule, field, data[field])
    db.session.commit()
    return Response.success(rule.to_dict(), 'AI规则已更新')


@b_followup_bp.route('/ai-rules/<int:rule_id>', methods=['DELETE'])
@login_required
def delete_ai_rule(rule_id):
    rule = BFollowUpAIRule.query.get_or_404(rule_id)
    db.session.delete(rule)
    db.session.commit()
    return Response.success({'id': rule_id}, 'AI规则已删除')


@b_followup_bp.route('/plans/candidates', methods=['GET'])
@login_required
def list_plan_candidates():
    """待制定随访计划患者列表"""
    search = request.args.get('search', '').strip()
    query = BPatient.query.filter(BPatient.status == 'active')
    if search:
        query = query.filter(db.or_(
            BPatient.name.like(f'%{search}%'),
            BPatient.phone.like(f'%{search}%'),
            BPatient.patient_code.like(f'%{search}%')
        ))
    patients = query.order_by(BPatient.created_at.desc()).limit(100).all()
    items = []
    for patient in patients:
        record = _latest_record(patient.id)
        report = _latest_report(patient.id, record.id if record else None)
        active_plan = BFollowUpPatientPlan.query.filter(
            BFollowUpPatientPlan.patient_id == patient.id,
            BFollowUpPatientPlan.status.in_(['draft', 'active', 'paused'])
        ).order_by(BFollowUpPatientPlan.created_at.desc()).first()
        data = patient.to_dict()
        data['latest_record_id'] = record.id if record else None
        data['latest_report_id'] = report.id if report else None
        data['risk_level'] = report.risk_level if report else None
        data['followup_plan'] = active_plan.to_dict() if active_plan else None
        data['plan_status'] = active_plan.status if active_plan else 'none'
        items.append(data)
    return Response.success(items)


@b_followup_bp.route('/plans/recommend', methods=['POST'])
@login_required
def recommend_patient_plan():
    """根据患者画像推荐随访模板、知识条目和节点内容"""
    data = request.json or {}
    patient_id = data.get('patient_id')
    if not patient_id:
        return Response.error('patient_id 必填', 400)
    patient = BPatient.query.get(patient_id)
    if not patient:
        return Response.error('患者不存在', 404)

    record = BHealthRecord.query.get(data.get('record_id')) if data.get('record_id') else _latest_record(patient.id)
    report = BReport.query.get(data.get('report_id')) if data.get('report_id') else _latest_report(patient.id, record.id if record else None)
    nodule_type = data.get('nodule_type') or patient.nodule_type
    risk_level = data.get('risk_level') or (report.risk_level if report else None)

    template = None
    if data.get('template_id'):
        template = BFollowUpPlanTemplate.query.get(data.get('template_id'))
    if not template:
        query = BFollowUpPlanTemplate.query.filter(BFollowUpPlanTemplate.status == 'active')
        if nodule_type:
            query = query.filter(db.or_(
                BFollowUpPlanTemplate.nodule_type == nodule_type,
                BFollowUpPlanTemplate.nodule_type.is_(None),
                BFollowUpPlanTemplate.nodule_type == ''
            ))
        if risk_level:
            query = query.filter(db.or_(
                BFollowUpPlanTemplate.risk_level.in_(_risk_aliases(risk_level)),
                BFollowUpPlanTemplate.risk_level.is_(None),
                BFollowUpPlanTemplate.risk_level == ''
            ))
        template = query.order_by(BFollowUpPlanTemplate.updated_at.desc()).first()
    if not template:
        return Response.error('暂无可推荐的随访模板，请先在随访知识库与模板页面导入或创建模板', 404)

    knowledge = _match_knowledge(
        patient=patient,
        nodule_type=nodule_type,
        risk_level=risk_level,
        categories=[
            'breast_nodule', 'lung_nodule', 'thyroid_nodule',
            'diet', 'exercise', 'psych', 'review', 'script', 'ai_rule', 'other'
        ],
        limit=30
    )
    nodes = template.nodes.filter(BFollowUpPlanNode.is_active.is_(True)).order_by(
        BFollowUpPlanNode.day_offset.asc(), BFollowUpPlanNode.sort_order.asc()
    ).all()

    node_payload = []
    knowledge_by_id = {item.id: item for item in knowledge}
    for node in nodes:
        node_items = []
        ids = node.knowledge_item_ids if isinstance(node.knowledge_item_ids, list) else []
        for item_id in ids:
            item = knowledge_by_id.get(item_id) or BFollowUpKnowledgeItem.query.get(item_id)
            if item:
                node_items.append(item.to_dict())
        if not node_items:
            node_items = [item.to_dict() for item in _match_knowledge(
                patient=patient,
                nodule_type=nodule_type,
                risk_level=risk_level,
                task_type=node.task_type,
                limit=5
            )]
        node_payload.append({**node.to_dict(), 'matched_knowledge': node_items})

    return Response.success({
        'patient': patient.to_dict(),
        'record_id': record.id if record else None,
        'report_id': report.id if report else None,
        'risk_level': risk_level,
        'nodule_type': nodule_type,
        'template': template.to_dict(include_nodes=False) if template else None,
        'nodes': node_payload,
        'knowledge': [item.to_dict() for item in knowledge],
        'settings': {
            'cycle_days': template.cycle_days if template else 90,
            'channel': template.default_channel if template else 'wecom',
            'reminder_strategy': template.default_reminder_strategy if template else '每日固定时间提醒；未打卡继续提醒；餐饮图片自动分析',
        }
    })


@b_followup_bp.route('/patient-plans', methods=['GET'])
@login_required
def list_patient_plans():
    patient_id = request.args.get('patient_id', type=int)
    status = request.args.get('status', '').strip()
    query = BFollowUpPatientPlan.query
    if patient_id:
        query = query.filter(BFollowUpPatientPlan.patient_id == patient_id)
    if status:
        query = query.filter(BFollowUpPatientPlan.status == status)
    items = query.order_by(BFollowUpPatientPlan.created_at.desc()).limit(100).all()
    return Response.success([item.to_dict() for item in items])


@b_followup_bp.route('/patient-plans', methods=['POST'])
@login_required
def create_patient_plan():
    """保存患者个性化随访计划"""
    data = request.json or {}
    patient_id = data.get('patient_id')
    if not patient_id:
        return Response.error('patient_id 必填', 400)
    patient = BPatient.query.get(patient_id)
    if not patient:
        return Response.error('患者不存在', 404)

    template = BFollowUpPlanTemplate.query.get(data.get('template_id')) if data.get('template_id') else None
    record = BHealthRecord.query.get(data.get('record_id')) if data.get('record_id') else _latest_record(patient.id)
    report = BReport.query.get(data.get('report_id')) if data.get('report_id') else _latest_report(patient.id, record.id if record else None)
    settings = data.get('settings') or {}
    plan_content = data.get('plan_content') or {}
    if not plan_content and template:
        plan_content = {
            'template': template.to_dict(include_nodes=False),
            'nodes': [node.to_dict() for node in template.nodes.filter(BFollowUpPlanNode.is_active.is_(True)).all()]
        }

    plan = BFollowUpPatientPlan(
        plan_code=data.get('plan_code') or _plan_code(),
        patient_id=patient.id,
        record_id=record.id if record else None,
        report_id=report.id if report else None,
        template_id=template.id if template else data.get('template_id'),
        manager_id=data.get('manager_id') or g.user_id,
        name=data.get('name') or (template.name if template else f'{patient.name}健康管理任务计划'),
        status=data.get('status') or 'draft',
        nodule_type=data.get('nodule_type') or patient.nodule_type,
        risk_level=data.get('risk_level') or (report.risk_level if report else None),
        cycle_days=settings.get('cycle_days') or data.get('cycle_days') or (template.cycle_days if template else 90),
        channel=settings.get('channel') or data.get('channel') or (template.default_channel if template else 'wecom'),
        reminder_strategy=settings.get('reminder_strategy') or data.get('reminder_strategy') or (template.default_reminder_strategy if template else ''),
        start_at=_parse_datetime(data.get('start_at')) or datetime.now(),
        plan_content=plan_content,
        selected_knowledge_ids=data.get('selected_knowledge_ids') or [],
    )
    db.session.add(plan)
    db.session.commit()
    return Response.success(plan.to_dict(), '患者随访计划已保存', 201)


def _node_send_datetime(start_at, node):
    base = start_at or datetime.now()
    day_offset = max(int(node.get('day_offset') or 1), 1)
    send_time = str(node.get('send_time') or '09:00')
    hour, minute = 9, 0
    try:
        parts = send_time.split(':')
        hour = int(parts[0])
        minute = int(parts[1]) if len(parts) > 1 else 0
    except Exception:
        pass
    target = base + timedelta(days=day_offset - 1)
    return target.replace(hour=hour, minute=minute, second=0, microsecond=0)


@b_followup_bp.route('/patient-plans/<int:plan_id>/activate', methods=['POST'])
@login_required
def activate_patient_plan(plan_id):
    """激活患者计划，并按计划节点生成实际随访任务"""
    plan = BFollowUpPatientPlan.query.get_or_404(plan_id)
    data = request.json or {}
    old_status = plan.status
    now = datetime.now()
    if data.get('start_at'):
        plan.start_at = _parse_datetime(data.get('start_at'))
    if not plan.start_at:
        plan.start_at = now
    plan.status = 'active'
    plan.activated_at = now
    plan.updated_at = now

    content = plan.plan_content or {}
    nodes = content.get('nodes') if isinstance(content, dict) else []
    if not nodes and plan.template:
        nodes = [node.to_dict() for node in plan.template.nodes.filter(BFollowUpPlanNode.is_active.is_(True)).all()]
    created_tasks = []
    for node in nodes or []:
        if not node.get('is_active', True):
            continue
        scheduled = _node_send_datetime(plan.start_at, node)
        task = BFollowUpTask(
            task_code=_task_code(),
            patient_id=plan.patient_id,
            record_id=plan.record_id,
            report_id=plan.report_id,
            manager_id=plan.manager_id,
            status='scheduled' if scheduled > now else 'pending',
            priority='high' if _normalize_risk_for_match(plan.risk_level) == 'high' else 'normal',
            risk_level=plan.risk_level,
            nodule_type=plan.nodule_type,
            source='patient_plan',
            title=f"{plan.patient.name if plan.patient else '用户'} · {node.get('name') or '健康管理任务'}",
            plan_name=plan.name,
            plan_day=node.get('day_offset') or 1,
            due_at=scheduled,
            scheduled_send_at=scheduled,
            channel=plan.channel,
            channel_recipient=plan.patient.wechat_id if plan.patient else None,
            ai_enabled=True,
            task_payload={
                'patient_plan_id': plan.id,
                'node': node,
                'reminder_strategy': plan.reminder_strategy,
            }
        )
        db.session.add(task)
        db.session.flush()
        _event(
            task,
            'task_created_from_patient_plan',
            f"由任务计划生成：{node.get('name') or '任务节点'}",
            to_status=task.status,
            actor_type='system',
            payload={'patient_plan_id': plan.id, 'node_code': node.get('node_code')}
        )
        created_tasks.append(task)
    db.session.commit()
    return Response.success({
        'plan': plan.to_dict(),
        'from_status': old_status,
        'tasks': [task.to_dict(include_events=True) for task in created_tasks]
    }, '任务计划已激活，提醒与打卡任务已生成')


@b_followup_bp.route('/checkins', methods=['POST'])
@login_required
def create_checkin():
    """记录患者打卡；图片识别/饮食点评后续在这里扩展"""
    data = request.json or {}
    patient_id = data.get('patient_id')
    if not patient_id:
        return Response.error('patient_id 必填', 400)
    patient = BPatient.query.get(patient_id)
    if not patient:
        return Response.error('患者不存在', 404)
    checkin_type = data.get('checkin_type') or 'general'
    ai_result = data.get('ai_result') or {}
    abnormal_flag = data.get('abnormal_flag', False)
    abnormal_reason = data.get('abnormal_reason')
    if data.get('analyze', True) and checkin_type == 'diet_checkin':
        ai_result = followup_checkin_service.analyze_diet(
            content_text=data.get('content_text') or '',
            image_urls=data.get('image_urls') or [],
            structured_data=data.get('structured_data') or {}
        )
        abnormal_flag = bool(ai_result.get('abnormal'))
        abnormal_reason = ai_result.get('summary') if abnormal_flag else abnormal_reason
    checkin_status = data.get('status') or ('alert' if abnormal_flag else 'analyzed' if ai_result else 'submitted')

    checkin = BFollowUpCheckin(
        task_id=data.get('task_id'),
        patient_plan_id=data.get('patient_plan_id'),
        patient_id=patient.id,
        checkin_type=checkin_type,
        content_text=data.get('content_text'),
        image_urls=data.get('image_urls') or [],
        structured_data=data.get('structured_data') or {},
        ai_result=ai_result,
        status=checkin_status,
        abnormal_flag=abnormal_flag,
        abnormal_reason=abnormal_reason
    )
    db.session.add(checkin)
    if checkin.task_id and abnormal_flag:
        task = BFollowUpTask.query.get(checkin.task_id)
        if task:
            old_status = task.status
            task.status = 'manual_processing'
            task.abnormal_flag = True
            task.abnormal_reason = abnormal_reason
            task.handoff_to = 'health_manager'
            task.ai_summary = ai_result.get('summary')
            _event(
                task,
                'checkin_alert',
                abnormal_reason or '打卡触发异常提醒',
                from_status=old_status,
                to_status=task.status,
                actor_type='ai',
                payload={'checkin_type': checkin_type, 'ai_result': ai_result}
            )
    db.session.commit()
    return Response.success(checkin.to_dict(), '打卡记录已保存', 201)


@b_followup_bp.route('/tasks', methods=['GET'])
@login_required
def list_tasks():
    """随访任务列表"""
    patient_id = request.args.get('patient_id', type=int)
    report_id = request.args.get('report_id', type=int)
    status = request.args.get('status', '').strip()
    search = request.args.get('search', '').strip()
    risk_level = request.args.get('risk_level', '').strip()
    channel = request.args.get('channel', '').strip()
    source = request.args.get('source', '').strip()
    due = request.args.get('due', '').strip()  # today/overdue
    page = request.args.get('page', 1, type=int)
    per_page = min(request.args.get('per_page', 20, type=int), 100)

    query = BFollowUpTask.query.join(BPatient)
    if patient_id:
        query = query.filter(BFollowUpTask.patient_id == patient_id)
    if report_id:
        query = query.filter(BFollowUpTask.report_id == report_id)
    if status and status != 'all':
        query = query.filter(BFollowUpTask.status == status)
    if risk_level and risk_level != 'all':
        query = query.filter(BFollowUpTask.risk_level == risk_level)
    if channel and channel != 'all':
        query = query.filter(BFollowUpTask.channel == channel)
    if source and source != 'all':
        query = query.filter(BFollowUpTask.source == source)
    if search:
        query = query.filter(db.or_(
            BPatient.name.like(f'%{search}%'),
            BPatient.phone.like(f'%{search}%'),
            BFollowUpTask.task_code.like(f'%{search}%')
        ))
    now = datetime.now()
    if due == 'today':
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        end = start + timedelta(days=1)
        query = query.filter(BFollowUpTask.due_at >= start, BFollowUpTask.due_at < end)
    elif due == 'overdue':
        query = query.filter(
            BFollowUpTask.due_at < now,
            BFollowUpTask.status.notin_(['completed', 'cancelled'])
        )

    query = query.order_by(BFollowUpTask.due_at.asc().nullslast(), BFollowUpTask.created_at.desc())
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    return Response.paginate(
        _serialize_task_query(pagination.items),
        pagination.total,
        page,
        per_page
    )


@b_followup_bp.route('/tasks/metrics', methods=['GET'])
@login_required
def task_metrics():
    """随访任务统计"""
    now = datetime.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    today_end = today_start + timedelta(days=1)
    base = BFollowUpTask.query
    data = {
        'today_due': base.filter(BFollowUpTask.due_at >= today_start, BFollowUpTask.due_at < today_end).count(),
        'pending': base.filter(BFollowUpTask.status.in_(['pending', 'scheduled'])).count(),
        'sent': base.filter(BFollowUpTask.status.in_(['sent', 'replied'])).count(),
        'alert': base.filter(BFollowUpTask.status.in_(['alert', 'manual_processing'])).count(),
        'completed': base.filter(BFollowUpTask.status == 'completed').count(),
        'overdue': base.filter(BFollowUpTask.due_at < now, BFollowUpTask.status.notin_(['completed', 'cancelled'])).count(),
        'ai_enabled': base.filter(BFollowUpTask.ai_enabled.is_(True)).count(),
    }
    total_closed_base = max(data['completed'] + data['alert'] + data['sent'] + data['pending'], 1)
    data['completion_rate'] = round(data['completed'] / total_closed_base * 100)
    return Response.success(data)


@b_followup_bp.route('/tasks/<int:task_id>', methods=['GET'])
@login_required
def get_task(task_id):
    task = BFollowUpTask.query.get_or_404(task_id)
    return Response.success(task.to_dict(include_messages=True, include_events=True))


@b_followup_bp.route('/tasks/<int:task_id>/checkins', methods=['GET'])
@login_required
def list_task_checkins(task_id):
    """查询单个随访任务的患者打卡记录"""
    task = BFollowUpTask.query.get_or_404(task_id)
    items = BFollowUpCheckin.query.filter_by(task_id=task.id).order_by(BFollowUpCheckin.submitted_at.desc()).all()
    return Response.success([item.to_dict() for item in items])


@b_followup_bp.route('/tasks', methods=['POST'])
@login_required
def create_task():
    """创建随访任务"""
    data = request.json or {}
    patient_id = data.get('patient_id')
    if not patient_id:
        return Response.error('patient_id 必填', 400)

    patient = BPatient.query.get(patient_id)
    if not patient:
        return Response.error('患者不存在', 404)

    record_id = data.get('record_id')
    report_id = data.get('report_id')
    record = BHealthRecord.query.get(record_id) if record_id else _latest_record(patient.id)
    report = BReport.query.get(report_id) if report_id else _latest_report(patient.id, record.id if record else None)
    due_at = _parse_datetime(data.get('due_at')) or (datetime.now() + timedelta(days=1)).replace(hour=9, minute=0, second=0, microsecond=0)
    scheduled_send_at = _parse_datetime(data.get('scheduled_send_at')) or due_at

    risk_level = data.get('risk_level') or (report.risk_level if report else None)
    nodule_type = data.get('nodule_type') or patient.nodule_type

    task = BFollowUpTask(
        task_code=_task_code(),
        patient_id=patient.id,
        record_id=record.id if record else None,
        report_id=report.id if report else None,
        manager_id=data.get('manager_id') or g.user_id,
        status=data.get('status') or 'pending',
        priority=data.get('priority') or ('high' if risk_level in ('高危', '高风险', 'high') else 'normal'),
        risk_level=risk_level,
        nodule_type=nodule_type,
        source=data.get('source') or 'manual',
        title=data.get('title') or f'{patient.name}随访任务',
        plan_name=data.get('plan_name') or '结节健康管理随访计划',
        plan_day=data.get('plan_day') or 1,
        due_at=due_at,
        scheduled_send_at=scheduled_send_at,
        channel=data.get('channel') or 'manual',
        channel_recipient=data.get('channel_recipient') or _recipient_for_channel(patient, data.get('channel') or 'manual'),
        ai_enabled=data.get('ai_enabled', True),
        task_payload=data.get('task_payload') or {}
    )
    db.session.add(task)
    db.session.flush()
    _event(task, 'task_created', '随访任务已创建', to_status=task.status, actor_type='staff')
    db.session.commit()
    return Response.success(task.to_dict(include_events=True), '随访任务创建成功', 201)


@b_followup_bp.route('/tasks/from-report/<int:report_id>', methods=['POST'])
@login_required
def create_task_from_report(report_id):
    """手动从已审核报告生成首次随访任务。"""
    report = BReport.query.get(report_id)
    if not report:
        return Response.error('报告不存在', 404)
    if report.status != 'finalized':
        return Response.error('报告尚未审核通过，不能生成报告后随访任务', 409)

    existing = _open_report_followup_task(report.id)
    if existing:
        return Response.error('该报告已存在随访任务，请勿重复生成', 409)

    patient = BPatient.query.get(report.patient_id)
    if not patient:
        return Response.error('患者不存在', 404)

    risk = report.risk_level or '未评估'
    days = _first_followup_days(risk)
    due_at = datetime.now() + timedelta(days=days)
    risk_key = _normalize_risk_for_match(risk)
    node = {
        'name': '报告后首次随访',
        'task_type': 'daily_checkin',
        'patient_action': 'reply_text',
        'ai_action': 'reply',
        'message_template': '您好，您的健康管理报告已完成审核。请完成本次报告后随访打卡，健康管理师会根据您的反馈继续跟进。',
        'checkin_schema': {
            'fields': ['睡眠', '饮食', '运动', '情绪', '症状变化', '备注'],
            'source': 'report_first_followup'
        }
    }
    task = BFollowUpTask(
        task_code=_task_code(),
        patient_id=patient.id,
        record_id=report.record_id,
        report_id=report.id,
        manager_id=g.user_id,
        status='pending',
        priority='high' if risk_key == 'high' else 'normal',
        risk_level=report.risk_level,
        nodule_type=patient.nodule_type,
        source='report',
        title=f'{patient.name}报告后首次随访',
        plan_name='报告审核后首次随访',
        plan_day=1,
        due_at=due_at.replace(hour=9, minute=0, second=0, microsecond=0),
        scheduled_send_at=datetime.now(),
        channel='manual',
        channel_recipient=_recipient_for_channel(patient, 'manual'),
        ai_enabled=True,
        task_payload={
            'report_id': report.id,
            'report_code': report.report_code,
            'report_summary': report.report_summary,
            'risk_level': report.risk_level,
            'first_followup_days': days,
            'rule': 'manual_report_first_followup',
            'node': node,
        }
    )
    db.session.add(task)
    db.session.flush()
    _event(task, 'task_created_from_report', '已根据报告生成随访任务', to_status='pending', actor_type='staff')
    db.session.commit()
    return Response.success(task.to_dict(include_events=True), '报告随访任务创建成功', 201)


@b_followup_bp.route('/tasks/<int:task_id>/send', methods=['POST'])
@login_required
def send_task(task_id):
    """生成AI话术并按渠道发送/记录；无企微时生成公开打卡链接。"""
    task = BFollowUpTask.query.get_or_404(task_id)
    data = request.json or {}
    content = (data.get('content') or '').strip()
    ai_result = None
    if not content:
        ai_result = followup_ai_service.build_outbound_message(task)
        content = ai_result.content

    send_result = _dispatch_followup_message(task, content)
    status = send_result.get('status') or ('sent' if send_result.get('ok') else 'failed')
    content = send_result.get('content') or content
    now = datetime.now()
    message = BFollowUpMessage(
        task_id=task.id,
        patient_id=task.patient_id,
        direction='outbound',
        sender_type='ai' if ai_result else 'staff',
        channel=task.channel,
        content_type='text',
        content=content,
        ai_intent=ai_result.intent if ai_result else 'manual_message',
        risk_signal=ai_result.risk_signal if ai_result else '',
        requires_manual_review=False,
        send_status=status,
        provider_message_id=send_result.get('provider_message_id'),
        provider_payload=send_result.get('payload'),
        provider_response=send_result.get('response') or send_result,
        error_message='' if send_result.get('ok') else send_result.get('reason') or str(send_result),
        sent_at=now if send_result.get('ok') else None,
    )
    db.session.add(message)
    db.session.flush()

    old_status = task.status
    task.last_message_id = message.id
    task.ai_summary = ai_result.summary if ai_result else '人工发送随访消息'
    if send_result.get('dry_run') and send_result.get('checkin_url'):
        task.ai_summary = f"{task.ai_summary}；公开打卡链接已生成"
    task.status = 'sent' if send_result.get('ok') else 'failed'
    task.updated_at = now
    _event(
        task,
        'message_sent' if send_result.get('ok') else 'message_failed',
        task.ai_summary,
        from_status=old_status,
        to_status=task.status,
        actor_type='ai' if ai_result else 'staff',
        payload=send_result
    )
    db.session.commit()
    return Response.success(task.to_dict(include_messages=True, include_events=True), '随访消息已处理')


@b_followup_bp.route('/tasks/<int:task_id>/reply', methods=['POST'])
@login_required
def record_reply(task_id):
    """记录患者回复并进行AI预警判定"""
    task = BFollowUpTask.query.get_or_404(task_id)
    data = request.json or {}
    content = (data.get('content') or '').strip()
    if not content:
        return Response.error('回复内容不能为空', 400)

    ai_result = followup_ai_service.analyze_patient_reply(content, task)
    now = datetime.now()
    inbound = BFollowUpMessage(
        task_id=task.id,
        patient_id=task.patient_id,
        direction='inbound',
        sender_type='patient',
        channel=data.get('channel') or task.channel,
        content_type='text',
        content=content,
        ai_intent=ai_result.intent,
        risk_signal=ai_result.risk_signal,
        requires_manual_review=ai_result.requires_manual_review,
        send_status='replied',
        provider_payload=data.get('provider_payload') or {},
        received_at=now,
    )
    db.session.add(inbound)
    db.session.flush()

    old_status = task.status
    task.last_message_id = inbound.id
    task.abnormal_flag = ai_result.abnormal or task.abnormal_flag
    task.abnormal_reason = ai_result.summary if ai_result.requires_manual_review else task.abnormal_reason
    task.handoff_to = ai_result.handoff_to or task.handoff_to
    task.ai_summary = ai_result.summary
    if ai_result.abnormal:
        task.status = 'alert'
    elif ai_result.requires_manual_review:
        task.status = 'manual_processing'
    else:
        task.status = 'replied'
    task.updated_at = now

    _event(
        task,
        'patient_reply_analyzed',
        ai_result.summary,
        from_status=old_status,
        to_status=task.status,
        actor_type='ai',
        payload=ai_result.to_dict()
    )
    db.session.commit()
    return Response.success(task.to_dict(include_messages=True, include_events=True), '患者回复已记录')


@b_followup_bp.route('/tasks/<int:task_id>/complete', methods=['POST'])
@login_required
def complete_task(task_id):
    task = BFollowUpTask.query.get_or_404(task_id)
    data = request.json or {}
    old_status = task.status
    task.status = 'completed'
    task.completed_at = datetime.now()
    task.updated_at = task.completed_at
    _event(
        task,
        'task_completed',
        data.get('summary') or '随访任务已闭环',
        from_status=old_status,
        to_status='completed',
        actor_type='staff',
        payload=data
    )
    db.session.commit()
    return Response.success(task.to_dict(include_events=True), '随访任务已完成')


@b_followup_bp.route('/tasks/<int:task_id>/manual-action', methods=['POST'])
@login_required
def manual_action_task(task_id):
    """健康管理师人工处理异常/待人工任务，并写入事件留痕。"""
    task = BFollowUpTask.query.get_or_404(task_id)
    data = request.json or {}
    action = data.get('action') or 'manual_followed'
    note = (data.get('note') or '').strip()
    old_status = task.status
    now = datetime.now()

    if action == 'doctor_handoff':
        task.status = 'manual_processing'
        task.handoff_to = 'doctor'
        summary = note or '已转医生处理'
    elif action == 'manual_followed':
        task.status = 'replied'
        task.handoff_to = 'health_manager'
        task.abnormal_flag = False
        summary = note or '健康管理师已人工跟进'
    elif action == 'close_alert':
        task.status = 'completed'
        task.completed_at = now
        task.abnormal_flag = False
        summary = note or '预警已复核并关闭'
    else:
        return Response.error('不支持的人工处理动作', 400)

    task.ai_summary = summary
    task.abnormal_reason = note or task.abnormal_reason
    task.updated_at = now
    _event(
        task,
        f'manual_{action}',
        summary,
        from_status=old_status,
        to_status=task.status,
        actor_type='staff',
        payload={'action': action, 'note': note}
    )
    db.session.commit()
    return Response.success(task.to_dict(include_messages=True, include_events=True), '人工处理已记录')


@b_followup_bp.route('/checkins/<int:checkin_id>/review', methods=['POST'])
@login_required
def review_checkin(checkin_id):
    """复核患者打卡记录。"""
    checkin = BFollowUpCheckin.query.get_or_404(checkin_id)
    data = request.json or {}
    checkin.status = data.get('status') or 'closed'
    checkin.reviewed_by = g.user_id
    checkin.reviewed_at = datetime.now()

    task = BFollowUpTask.query.get(checkin.task_id) if checkin.task_id else None
    if task:
        old_status = task.status
        if checkin.abnormal_flag and data.get('task_status'):
            task.status = data.get('task_status')
        task.updated_at = datetime.now()
        _event(
            task,
            'checkin_reviewed',
            data.get('note') or '患者打卡已复核',
            from_status=old_status,
            to_status=task.status,
            actor_type='staff',
            payload={'checkin_id': checkin.id, 'status': checkin.status, 'note': data.get('note')}
        )
    db.session.commit()
    return Response.success(checkin.to_dict(), '打卡已复核')


def _parse_xml_payload(raw):
    if not raw:
        return {}
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return {'raw': raw}
    return {child.tag: child.text for child in root}


def _read_wecom_payload():
    data = request.get_json(silent=True)
    if isinstance(data, dict):
        return data
    form = request.form.to_dict()
    if form:
        return form
    return _parse_xml_payload(request.get_data(as_text=True))


def _wecom_aes_key():
    key = current_app.config.get('WECOM_CALLBACK_AES_KEY', '')
    if not key:
        raise ValueError('未配置 WECOM_CALLBACK_AES_KEY')
    try:
        aes_key = base64.b64decode(f'{key}=')
    except Exception as exc:
        raise ValueError('WECOM_CALLBACK_AES_KEY 格式错误') from exc
    if len(aes_key) != 32:
        raise ValueError('WECOM_CALLBACK_AES_KEY 解码后必须为 32 字节')
    return aes_key


def _pkcs7_unpad(data):
    if not data:
        raise ValueError('企业微信回调密文为空')
    pad = data[-1]
    if pad < 1 or pad > 32:
        raise ValueError('企业微信回调 padding 无效')
    if data[-pad:] != bytes([pad]) * pad:
        raise ValueError('企业微信回调 padding 校验失败')
    return data[:-pad]


def _decrypt_wecom_encrypt(encrypted_text):
    aes_key = _wecom_aes_key()
    try:
        encrypted = base64.b64decode(encrypted_text)
    except Exception as exc:
        raise ValueError('企业微信回调 Encrypt 不是有效 base64') from exc
    cipher = Cipher(algorithms.AES(aes_key), modes.CBC(aes_key[:16]), backend=default_backend())
    decryptor = cipher.decryptor()
    plain = decryptor.update(encrypted) + decryptor.finalize()
    plain = _pkcs7_unpad(plain)
    if len(plain) < 20:
        raise ValueError('企业微信回调明文长度无效')

    msg_len = struct.unpack('!I', plain[16:20])[0]
    xml_start = 20
    xml_end = xml_start + msg_len
    xml_bytes = plain[xml_start:xml_end]
    receive_id = plain[xml_end:].decode('utf-8', errors='ignore')
    expected_receive_id = current_app.config.get('WECOM_CORP_ID', '')
    if expected_receive_id and receive_id and receive_id != expected_receive_id:
        raise ValueError('企业微信回调 CorpID 校验失败')
    return xml_bytes.decode('utf-8'), receive_id


def _wecom_signature_ok(encrypted_text=''):
    """
    企业微信签名校验框架。

    默认开发环境允许明文/mock 回调；生产可配置 WECOM_CALLBACK_VERIFY_SIGNATURE=True
    后强制校验 token/timestamp/nonce/msg_signature。
    """
    verify_required = current_app.config.get('WECOM_CALLBACK_VERIFY_SIGNATURE', False)
    token = current_app.config.get('WECOM_CALLBACK_TOKEN', '')
    signature = request.args.get('msg_signature') or request.args.get('signature') or ''
    timestamp = request.args.get('timestamp') or ''
    nonce = request.args.get('nonce') or ''
    if not verify_required and not (token and signature):
        return True
    if not token or not signature or not timestamp or not nonce:
        return False
    parts = [token, timestamp, nonce]
    if encrypted_text:
        parts.append(encrypted_text)
    digest = hashlib.sha1(''.join(sorted(parts)).encode('utf-8')).hexdigest()
    return digest == signature


def _payload_value(payload, *keys):
    lower = {str(k).lower(): v for k, v in (payload or {}).items()}
    for key in keys:
        if key in payload and payload.get(key) not in (None, ''):
            return payload.get(key)
        value = lower.get(str(key).lower())
        if value not in (None, ''):
            return value
    return ''


def _wecom_external_userid(payload):
    return _payload_value(
        payload,
        'ExternalUserID',
        'ExternalUserId',
        'external_userid',
        'externalUserId',
        'FromUserName',
        'from_user_name',
        'userid',
    )


def _wecom_message_type(payload):
    return _payload_value(payload, 'MsgType', 'msgtype', 'msg_type', 'Event', 'event') or 'event'


def _wecom_message_content(payload):
    msg_type = _wecom_message_type(payload)
    if msg_type == 'text':
        return _payload_value(payload, 'Content', 'content') or ''
    if msg_type in ('image', 'voice', 'video', 'file'):
        media_id = _payload_value(payload, 'MediaId', 'media_id')
        pic_url = _payload_value(payload, 'PicUrl', 'pic_url', 'ImageUrl', 'image_url')
        return pic_url or media_id or f'收到{msg_type}消息'
    event = _payload_value(payload, 'Event', 'event')
    return event or f'收到企业微信{msg_type}回调'


def _wecom_media_info(payload):
    media_id = _payload_value(payload, 'MediaId', 'media_id')
    pic_url = _payload_value(payload, 'PicUrl', 'pic_url', 'ImageUrl', 'image_url')
    result = {
        'media_id': media_id,
        'pic_url': pic_url,
        'download': None,
        'image_urls': [pic_url] if pic_url else [],
    }
    if media_id:
        try:
            download = wecom_service.download_media(media_id)
        except Exception as exc:
            download = {'ok': False, 'status': 'failed', 'media_id': media_id, 'reason': str(exc)}
        result['download'] = download
        if download.get('url'):
            result['image_urls'].insert(0, download['url'])
    return result


def _find_wecom_task(patient, payload):
    task_id = _payload_value(payload, 'task_id', 'TaskId', 'taskId')
    if task_id:
        task = BFollowUpTask.query.filter_by(id=task_id, patient_id=patient.id).first()
        if task:
            return task

    active_statuses = ['sent', 'scheduled', 'pending', 'replied', 'manual_processing', 'alert']
    task = (
        BFollowUpTask.query
        .filter(BFollowUpTask.patient_id == patient.id)
        .filter(BFollowUpTask.channel == 'wecom')
        .filter(BFollowUpTask.status.in_(active_statuses))
        .order_by(BFollowUpTask.updated_at.desc(), BFollowUpTask.created_at.desc())
        .first()
    )
    if task:
        return task

    manager_id = _fallback_manager_id(patient)
    if not manager_id:
        return None

    now = datetime.now()
    task = BFollowUpTask(
        task_code=_task_code(),
        patient_id=patient.id,
        manager_id=manager_id,
        status='replied',
        priority='normal',
        risk_level=patient.risk_level if hasattr(patient, 'risk_level') else None,
        nodule_type=patient.nodule_type,
        source='wecom_callback',
        title=f'{patient.name}企微会话',
        plan_name='企业微信患者交互',
        plan_day=1,
        due_at=now,
        scheduled_send_at=now,
        channel='wecom',
        channel_recipient=patient.wecom_external_userid or patient.wecom_userid,
        ai_enabled=True,
        task_payload={'source': 'wecom_callback'}
    )
    db.session.add(task)
    db.session.flush()
    _event(task, 'task_created_from_wecom_callback', '患者企微消息自动创建会话任务', to_status=task.status, actor_type='wecom', payload=payload)
    return task


def _create_checkin_from_wecom_media(task, patient, payload, media_info, content):
    msg_type = _wecom_message_type(payload)
    if msg_type not in ('image', 'file'):
        return None, None

    image_urls = media_info.get('image_urls') or []
    structured_data = {
        'source': 'wecom_callback',
        'msg_type': msg_type,
        'media_id': media_info.get('media_id'),
        'pic_url': media_info.get('pic_url'),
        'download': media_info.get('download'),
        'raw_payload': payload,
    }
    checkin_type = 'diet_checkin' if msg_type == 'image' else 'image_checkin'
    ai_result = {}
    abnormal_flag = False
    abnormal_reason = None
    if msg_type == 'image':
        ai_result = followup_checkin_service.analyze_diet(content, image_urls, structured_data)
        abnormal_flag = bool(ai_result.get('abnormal'))
        abnormal_reason = ai_result.get('summary') if abnormal_flag else None

    checkin = BFollowUpCheckin(
        task_id=task.id,
        patient_plan_id=(task.task_payload or {}).get('patient_plan_id') if isinstance(task.task_payload, dict) else None,
        patient_id=patient.id,
        checkin_type=checkin_type,
        content_text=content or '患者通过企业微信上传图片',
        image_urls=image_urls,
        structured_data=structured_data,
        ai_result=ai_result,
        status='alert' if abnormal_flag else 'analyzed' if ai_result else 'submitted',
        abnormal_flag=abnormal_flag,
        abnormal_reason=abnormal_reason,
        submitted_at=datetime.now(),
    )
    db.session.add(checkin)
    db.session.flush()
    return checkin, ai_result


def _handle_wecom_callback():
    if request.method == 'GET':
        echostr = request.args.get('echostr', '')
        if not _wecom_signature_ok(echostr):
            return 'invalid signature', 403
        if request.args.get('msg_signature') and current_app.config.get('WECOM_CALLBACK_AES_KEY'):
            try:
                decrypted_echo, _receive_id = _decrypt_wecom_encrypt(echostr)
                return decrypted_echo or 'ok'
            except ValueError as exc:
                return str(exc), 400
        return echostr or 'ok'

    payload = _read_wecom_payload()
    encrypted = _payload_value(payload, 'Encrypt')
    if not _wecom_signature_ok(encrypted):
        return Response.error('企业微信回调签名校验失败', 403)
    if encrypted:
        try:
            decrypted_xml, receive_id = _decrypt_wecom_encrypt(encrypted)
            decrypted_payload = _parse_xml_payload(decrypted_xml)
            payload = {
                **decrypted_payload,
                '_encrypted': True,
                '_receive_id': receive_id,
                '_raw_encrypted_payload': payload,
            }
        except ValueError as exc:
            return Response.error(f'企业微信回调解密失败：{exc}', 400)

    external_userid = _wecom_external_userid(payload)
    if not external_userid:
        return Response.success({'received': True, 'linked': False, 'reason': 'missing_external_userid'}, '企业微信回调已接收，缺少 external_userid')

    patient = (
        BPatient.query
        .filter(
            (BPatient.wecom_external_userid == external_userid) |
            (BPatient.wecom_userid == external_userid)
        )
        .first()
    )
    if not patient:
        return Response.success({'received': True, 'linked': False, 'external_userid': external_userid}, '企业微信回调已接收，未找到绑定患者')

    task = _find_wecom_task(patient, payload)
    if not task:
        return Response.success({'received': True, 'linked': False, 'patient_id': patient.id, 'reason': 'missing_manager'}, '企业微信回调已接收，患者缺少负责人')

    now = datetime.now()
    msg_type = _wecom_message_type(payload)
    content = _wecom_message_content(payload)
    media_info = _wecom_media_info(payload) if msg_type in ('image', 'file') else {'image_urls': []}
    old_status = task.status
    ai_result = None
    analyze_types = {'text'}
    if msg_type in analyze_types and content:
        ai_result = followup_ai_service.analyze_patient_reply(content, task)

    message = BFollowUpMessage(
        task_id=task.id,
        patient_id=patient.id,
        direction='inbound',
        sender_type='patient',
        channel='wecom',
        content_type=msg_type,
        content=content or '收到企业微信回调',
        ai_intent=ai_result.intent if ai_result else msg_type,
        risk_signal=ai_result.risk_signal if ai_result else '',
        requires_manual_review=bool(ai_result.requires_manual_review) if ai_result else msg_type in ('image', 'file'),
        send_status='replied',
        provider_message_id=_payload_value(payload, 'MsgId', 'msgid', 'MsgID'),
        provider_payload={**payload, 'media_info': media_info} if media_info.get('media_id') or media_info.get('pic_url') else payload,
        received_at=now,
    )
    db.session.add(message)
    db.session.flush()

    checkin = None
    checkin_ai_result = None
    if msg_type in ('image', 'file'):
        checkin, checkin_ai_result = _create_checkin_from_wecom_media(task, patient, payload, media_info, content)

    task.last_message_id = message.id
    if ai_result:
        task.abnormal_flag = ai_result.abnormal or task.abnormal_flag
        task.abnormal_reason = ai_result.summary if ai_result.requires_manual_review else task.abnormal_reason
        task.handoff_to = ai_result.handoff_to or task.handoff_to
        task.ai_summary = ai_result.summary
        task.status = 'alert' if ai_result.abnormal else 'manual_processing' if ai_result.requires_manual_review else 'replied'
    elif checkin:
        task.abnormal_flag = bool(checkin.abnormal_flag) or task.abnormal_flag
        task.abnormal_reason = checkin.abnormal_reason or '患者上传图片/文件，待AI识别或人工点评'
        task.ai_summary = checkin_ai_result.get('summary') if checkin_ai_result else task.abnormal_reason
        task.status = 'alert' if checkin.abnormal_flag else 'manual_processing'
    else:
        task.status = 'replied'
        task.ai_summary = content or '收到企业微信回调'
    task.updated_at = now

    _event(
        task,
        'wecom_message_received',
        task.ai_summary or '收到企业微信患者消息',
        from_status=old_status,
        to_status=task.status,
        actor_type='wecom',
        payload={
            'external_userid': external_userid,
            'msg_type': msg_type,
            'message_id': message.id,
            'checkin_id': checkin.id if checkin else None,
            'media_info': media_info,
            'ai_result': ai_result.to_dict() if ai_result else checkin_ai_result,
        }
    )
    db.session.commit()
    return Response.success({
        'received': True,
        'linked': True,
        'patient_id': patient.id,
        'task_id': task.id,
        'message_id': message.id,
        'checkin_id': checkin.id if checkin else None,
        'status': task.status,
    }, '企业微信回调已记录')


@b_followup_bp.route('/wecom/callback', methods=['GET', 'POST'])
@wecom_callback_bp.route('/callback', methods=['GET', 'POST'])
def wecom_callback():
    """企业微信回调入口：支持 URL 验证、明文/mock 消息接入和患者归属。"""
    return _handle_wecom_callback()
